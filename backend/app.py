import io
import json
import re
import secrets
import smtplib
import string
import uuid
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from email.header import Header
from email.mime.text import MIMEText
from email.utils import formataddr
from pathlib import Path

from flask import Flask, jsonify, request, send_file, send_from_directory
from openpyxl import Workbook, load_workbook
from flask_cors import CORS
from werkzeug.security import check_password_hash, generate_password_hash

import config
from db import get_db, init_db

app = Flask(__name__)
CORS(app)
init_db()

UPLOAD_DIR = Path(__file__).parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
ALLOWED_IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "webp"}


def ensure_super_admin():
    username = (getattr(config, "SUPER_ADMIN_USERNAME", "") or "").strip()
    if not username:
        return
    password = getattr(config, "SUPER_ADMIN_PASSWORD", "") or "admin123456"
    conn = get_db()
    row = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
    if row:
        if row["role"] != "super" or not check_password_hash(row["password_hash"], password):
            conn.execute(
                "UPDATE users SET role = 'super', password_hash = ? WHERE id = ?",
                (generate_password_hash(password), row["id"]),
            )
    else:
        conn.execute(
            "INSERT INTO users (username, password_hash, token, role) VALUES (?, ?, ?, 'super')",
            (username, generate_password_hash(password), secrets.token_hex(16)),
        )
    conn.commit()
    conn.close()


ensure_super_admin()


def current_user():
    token = request.headers.get("Authorization", "")
    if not token:
        return None
    conn = get_db()
    row = conn.execute("SELECT * FROM users WHERE token = ?", (token,)).fetchone()
    conn.close()
    return row


def is_admin(user):
    return user is not None and user["role"] in ("admin", "super")


def is_super(user):
    return user is not None and user["role"] == "super"


USERNAME_RE = re.compile(r"[一-龥A-Za-z0-9]{2,20}")


def validate_username(username):
    if not USERNAME_RE.fullmatch(username):
        return "用户名需为 2-20 个汉字、字母或数字"
    return None


def validate_password(password):
    if len(password) < 6 or len(password) > 32:
        return "密码需为 6-32 个字符"
    if not re.search(r"[A-Za-z]", password) or not re.search(r"\d", password):
        return "密码需同时包含字母和数字"
    return None


def find_blocked_words(*texts):
    words = getattr(config, "BLOCKED_WORDS", []) or []
    hits = []
    for t in texts:
        if not t:
            continue
        low = str(t).lower()
        for w in words:
            w = (w or "").strip()
            if w and w.lower() in low and w not in hits:
                hits.append(w)
    return hits


def scan_questionnaire(qid):
    conn = get_db()
    q = conn.execute("SELECT title, description FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    texts = [q["title"], q["description"]] if q else []
    rows = conn.execute("SELECT title FROM questions WHERE questionnaire_id = ?", (qid,)).fetchall()
    texts += [r["title"] for r in rows]
    opts = conn.execute(
        "SELECT text FROM options WHERE question_id IN (SELECT id FROM questions WHERE questionnaire_id = ?)",
        (qid,),
    ).fetchall()
    texts += [r["text"] for r in opts]
    conn.close()
    return find_blocked_words(*texts)


def send_verification_email(to_addr, code):
    subject = "问卷系统 - 邮箱验证码"
    body = f"你的验证码是：{code}，10 分钟内有效。请勿泄露给他人。"
    msg = MIMEText(body, "plain", "utf-8")
    msg["Subject"] = Header(subject, "utf-8")
    msg["From"] = formataddr((config.SMTP_FROM, config.SMTP_USER))
    msg["To"] = to_addr

    server = smtplib.SMTP_SSL(config.SMTP_HOST, config.SMTP_PORT)
    server.login(config.SMTP_USER, config.SMTP_PASSWORD)
    server.sendmail(config.SMTP_USER, [to_addr], msg.as_string())
    server.quit()


# 猎人密钥字母表：去掉易混淆的 I / O / 0 / 1
HUNTER_KEY_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"


def generate_hunter_key():
    parts = [
        "".join(secrets.choice(HUNTER_KEY_ALPHABET) for _ in range(4))
        for _ in range(3)
    ]
    return "HT-" + "-".join(parts)


def send_hunter_key_email(to_addr, key):
    subject = "猎人协会 · 考前训练密钥"
    body = (
        "亲爱的准猎人：\n\n"
        "猎人协会已受理您的第 289 届猎人资格考试报名申请，\n"
        "并为您开通了「考前训练」秘密渠道。\n\n"
        f"您的专属密钥如下：\n\n    {key}\n\n"
        "使用方法：进入猎人协会首页，点击 H 徽章中央的红色菱形，\n"
        "在弹出的窗口中输入上述密钥即可进入训练。\n\n"
        "此密钥长期有效，请妥善保管，切勿泄露。\n\n"
        "—— 猎人协会 · 本部"
    )
    msg = MIMEText(body, "plain", "utf-8")
    msg["Subject"] = Header(subject, "utf-8")
    msg["From"] = formataddr((config.SMTP_FROM, config.SMTP_USER))
    msg["To"] = to_addr

    server = smtplib.SMTP_SSL(config.SMTP_HOST, config.SMTP_PORT)
    server.login(config.SMTP_USER, config.SMTP_PASSWORD)
    server.sendmail(config.SMTP_USER, [to_addr], msg.as_string())
    server.quit()


@app.route("/api/hunter/key", methods=["POST"])
def hunter_key():
    data = request.get_json(silent=True) or {}
    email = (data.get("email") or "").strip()
    classification = (data.get("classification") or "").strip()
    if not email or "@" not in email:
        return jsonify({"error": "邮箱格式不正确"}), 400

    if not config.SMTP_USER or not config.SMTP_PASSWORD:
        return jsonify({"error": "SMTP 未配置，请在 config.py 中填写邮箱和授权码"}), 500

    conn = get_db()
    row = conn.execute("SELECT * FROM hunter_keys WHERE email = ?", (email,)).fetchone()
    if row:
        key = row["key"]
    else:
        key = generate_hunter_key()
        conn.execute("INSERT INTO hunter_keys (email, key) VALUES (?, ?)", (email, key))

    app_row = conn.execute(
        "SELECT id FROM hunter_applications WHERE email = ?", (email,)
    ).fetchone()
    if app_row:
        conn.execute(
            "UPDATE hunter_applications SET classification = ? WHERE email = ?",
            (classification, email),
        )
    else:
        conn.execute(
            "INSERT INTO hunter_applications (email, classification) VALUES (?, ?)",
            (email, classification),
        )
    conn.commit()
    conn.close()

    try:
        send_hunter_key_email(email, key)
    except Exception as e:
        return jsonify({"error": f"邮件发送失败：{e}"}), 500

    return jsonify({"ok": True})


@app.route("/api/hunter/train/verify", methods=["POST"])
def hunter_train_verify():
    data = request.get_json(silent=True) or {}
    key = (data.get("key") or "").strip().upper()

    conn = get_db()
    row = conn.execute("SELECT * FROM hunter_keys WHERE key = ?", (key,)).fetchone()
    if row is None:
        conn.close()
        return jsonify({"error": "密钥无效，请检查后重试"}), 400

    app_row = conn.execute(
        "SELECT classification FROM hunter_applications WHERE email = ?",
        (row["email"],),
    ).fetchone()
    conn.close()

    return jsonify({
        "ok": True,
        "classification": app_row["classification"] if app_row else "",
    })


@app.route("/api/register", methods=["POST"])
def register():
    data = request.get_json(silent=True) or {}
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""

    if not username or not password:
        return jsonify({"error": "用户名和密码不能为空"}), 400

    err = validate_username(username)
    if err:
        return jsonify({"error": err}), 400
    err = validate_password(password)
    if err:
        return jsonify({"error": err}), 400

    super_name = (getattr(config, "SUPER_ADMIN_USERNAME", "") or "").strip()
    if super_name and username.lower() == super_name.lower():
        return jsonify({"error": "该用户名已被保留"}), 400

    conn = get_db()
    exists = conn.execute(
        "SELECT id FROM users WHERE LOWER(username) = LOWER(?)", (username,)
    ).fetchone()
    if exists:
        conn.close()
        return jsonify({"error": "用户名已存在"}), 400

    token = secrets.token_hex(16)
    conn.execute(
        "INSERT INTO users (username, password_hash, token) VALUES (?, ?, ?)",
        (username, generate_password_hash(password), token),
    )
    conn.commit()
    conn.close()

    return jsonify({"token": token, "username": username})


@app.route("/api/login", methods=["POST"])
def login():
    data = request.get_json(silent=True) or {}
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""

    conn = get_db()
    row = conn.execute("SELECT * FROM users WHERE LOWER(username) = LOWER(?)", (username,)).fetchone()
    conn.close()

    if row is None or not check_password_hash(row["password_hash"], password):
        return jsonify({"error": "用户名或密码错误"}), 400

    return jsonify({"token": row["token"], "username": row["username"]})


@app.route("/api/me", methods=["GET", "PUT"])
def me():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    if request.method == "PUT":
        data = request.get_json(silent=True) or {}
        conn = get_db()

        username = (data.get("username") or "").strip()
        if username and username != user["username"]:
            err = validate_username(username)
            if err:
                conn.close()
                return jsonify({"error": err}), 400
            exists = conn.execute(
                "SELECT id FROM users WHERE LOWER(username) = LOWER(?) AND id != ?",
                (username, user["id"]),
            ).fetchone()
            if exists:
                conn.close()
                return jsonify({"error": "昵称已被占用"}), 400
            conn.execute("UPDATE users SET username = ? WHERE id = ?", (username, user["id"]))

        avatar = data.get("avatar")
        if avatar is not None:
            conn.execute("UPDATE users SET avatar = ? WHERE id = ?", (avatar, user["id"]))

        email = data.get("email")
        if email is not None:
            conn.execute("UPDATE users SET email = ? WHERE id = ?", ((email or "").strip(), user["id"]))

        phone = data.get("phone")
        if phone is not None:
            conn.execute("UPDATE users SET phone = ? WHERE id = ?", ((phone or "").strip(), user["id"]))

        new_password = data.get("new_password")
        if new_password:
            err = validate_password(new_password)
            if err:
                conn.close()
                return jsonify({"error": err}), 400
            old_password = data.get("old_password") or ""
            if not check_password_hash(user["password_hash"], old_password):
                conn.close()
                return jsonify({"error": "旧密码错误"}), 400
            conn.execute(
                "UPDATE users SET password_hash = ? WHERE id = ?",
                (generate_password_hash(new_password), user["id"]),
            )

        conn.commit()
        row = conn.execute("SELECT * FROM users WHERE id = ?", (user["id"],)).fetchone()
        conn.close()
        return jsonify({
            "id": row["id"],
            "username": row["username"],
            "avatar": row["avatar"],
            "email": row["email"],
            "phone": row["phone"],
            "role": row["role"],
        })

    return jsonify({
        "id": user["id"],
        "username": user["username"],
        "avatar": user["avatar"],
        "email": user["email"],
        "phone": user["phone"],
        "role": user["role"],
    })


@app.route("/api/upload", methods=["POST"])
def upload_image():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "未选择文件"}), 400

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_IMAGE_EXT:
        return jsonify({"error": "仅支持图片格式"}), 400

    filename = f"{uuid.uuid4().hex}.{ext}"
    file.save(UPLOAD_DIR / filename)
    return jsonify({"url": f"/uploads/{filename}"})


@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(UPLOAD_DIR, filename)


EXAM_TEMPLATE_HEADERS = ["题型", "题干", "选项A", "选项B", "选项C", "选项D", "正确答案", "分值"]
EXAM_TEMPLATE_ROWS = [
    ["单选", "星露谷物语中，第一年春季可以种植的作物是？", "防风草", "草莓", "南瓜", "玉米", "A", 5],
    ["多选", "以下哪些是星露谷的节日？", "复活节", "月光水母节", "冬星节", "万圣节", "ABC", 10],
    ["判断", "星露谷中，钓鱼可以在任何季节进行。", "对", "错", "", "", "A", 5],
    ["填空", "星露谷中，冬天举办的大型节日叫____。", "", "", "", "", "冬星节", 5],
]

EXAM_TYPE_MAP = {
    "单选": "single", "单选题": "single",
    "多选": "multiple", "多选题": "multiple",
    "判断": "judge", "判断题": "judge",
    "填空": "fill", "填空题": "fill",
}

# DeepSeek AI（每个用户自己填 API Key，这里只放服务地址和模型名）
DEEPSEEK_BASE_URL = "https://api.deepseek.com"
DEEPSEEK_MODEL = "deepseek-chat"

EXAM_AI_SYSTEM = (
    "你是考试出卷助手。根据用户的描述生成一套考试题目，只输出一个 json 对象，不要输出任何解释或额外文字。\n"
    "json 结构必须严格为：\n"
    '{"questions": [{"type": "single|multiple|judge|fill", "title": "题干", "score": 整数分值, '
    '"options": [{"text": "选项文字", "is_correct": true或false}], "answer": "仅填空题填参考答案，多个用|分隔"}]}\n'
    "规则：\n"
    "1. type 只能是 single(单选)、multiple(多选)、judge(判断)、fill(填空) 之一。\n"
    "2. single/multiple 必须含 options 数组，每项有 text 和 is_correct 布尔值；single 只能有一个 is_correct=true，multiple 可有多个。\n"
    "3. judge 的 options 固定为 [{\"text\":\"对\",\"is_correct\":布尔},{\"text\":\"错\",\"is_correct\":布尔}]。\n"
    "4. fill 不需要 options，用 answer 字段填参考答案，多个答案用 | 分隔。\n"
    "5. score 为正整数，通常 5 分。题目数量遵循用户要求，默认 5 题。"
)

TEST_AI_SYSTEM = (
    "你是问卷出卷助手。根据用户的描述生成一套测试问卷（含题目和结果卡片），只输出一个 json 对象，不要输出任何解释或额外文字。\n"
    "json 结构必须严格为：\n"
    '{"questions": [{"type": "single|scale", "title": "题干", "options": [{"text": "选项文字", "score": 整数分值}]}], '
    '"result_cards": [{"min_score": 整数最低分, "max_score": 整数最高分, "text": "结果文案"}]}\n'
    "规则：\n"
    "1. type 只能是 single(单选题) 或 scale(量表题)。\n"
    "2. single 的每个选项 score 为整数分值，可正可负；scale 的 options 固定为 "
    "[{\"text\":\"完全不符合\",\"score\":1},{\"text\":\"不符合\",\"score\":2},{\"text\":\"中立\",\"score\":3},"
    "{\"text\":\"符合\",\"score\":4},{\"text\":\"完全符合\",\"score\":5}]。\n"
    "3. result_cards 的分数区间要覆盖所有可能得分、互不重叠，min_score<=max_score，text 是该分数段对应的结果文案。"
)

JUMP_AI_SYSTEM = (
    "你是问卷出卷助手。根据用户的描述生成一套跳转式测试问卷（每题为单选题，通过选项跳转到下一题或直接跳到某个结果），只输出一个 json 对象，不要输出任何解释或额外文字。\n"
    "json 结构必须严格为：\n"
    '{"questions": [{"type": "single", "title": "题干", "options": [{"text": "选项文字", "jump_to": "跳转目标"}]}], '
    '"result_cards": [{"label": "A|B|C...", "text": "结果文案"}]}\n'
    "规则：\n"
    "1. questions 只含 single 单选题。每个选项的 jump_to 表示选中该选项后跳转到的位置。\n"
    "2. jump_to 取值：'qN' 表示跳到第 N 题（N 为正整数，从 1 开始）；'rX' 表示直接跳到结果 X（X 为 result_cards 里的 label，从 A 开始的大写字母）。\n"
    "3. 只能向后跳：'qN' 的 N 必须大于当前题号，禁止跳回前面的题，避免循环。\n"
    "4. 最后一题的每个选项必须用 'rX' 跳到某个结果，不能再跳题。\n"
    "5. 每个结果（label）至少要被某个选项引用一次；除第 1 题外，每道题至少要被前面某题的选项通过 'qN' 引用到。\n"
    "6. 整体像一棵从第 1 题出发、逐层分叉、最终落到结果的决策树。\n"
    "7. result_cards 每个结果含 label（从 A 开始连续的大写字母）和 text（结果文案），数量通常 3~5 个。\n"
    "8. 题目数量遵循用户要求，默认 5 题。"
)


def sanitize_jump(questions, result_cards):
    labels = []
    for c in result_cards or []:
        if isinstance(c, dict):
            lab = (c.get("label") or "").strip()
            if lab:
                labels.append(lab)
    n = len(questions)
    for i, q in enumerate(questions):
        if not isinstance(q, dict):
            continue
        opts = q.get("options")
        if not isinstance(opts, list):
            continue
        for o in opts:
            if not isinstance(o, dict):
                continue
            jt = (o.get("jump_to") or "").strip()
            valid = ""
            if jt.startswith("q"):
                try:
                    num = int(jt[1:])
                except ValueError:
                    num = 0
                # 1-based forward jump: num-1 must be after current index i
                if 1 <= num <= n and num > i + 1:
                    valid = "q" + str(num)
            elif jt.startswith("r"):
                if jt[1:] in labels:
                    valid = "r" + jt[1:]
            if not valid and i == n - 1 and labels:
                valid = "r" + labels[0]
            o["jump_to"] = valid


def call_deepseek(api_key, system, prompt):
    url = DEEPSEEK_BASE_URL.rstrip("/") + "/chat/completions"
    payload = {
        "model": DEEPSEEK_MODEL,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.7,
        "response_format": {"type": "json_object"},
    }
    data = json.dumps(payload).encode("utf-8")
    req = Request(url, data=data, method="POST")
    req.add_header("Content-Type", "application/json")
    req.add_header("Authorization", "Bearer " + api_key)
    with urlopen(req, timeout=60) as resp:
        body = json.loads(resp.read().decode("utf-8"))
    content = body["choices"][0]["message"]["content"]
    return json.loads(content)


@app.route("/api/exam/template", methods=["GET"])
def exam_template():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    wb = Workbook()
    ws = wb.active
    ws.title = "考试试卷"
    ws.append(EXAM_TEMPLATE_HEADERS)
    for row in EXAM_TEMPLATE_ROWS:
        ws.append(row)

    widths = [10, 40, 14, 14, 14, 14, 12, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="考试试卷模板.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/exam/parse", methods=["POST"])
def exam_parse():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "未选择文件"}), 400

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ("xlsx", "xlsm"):
        return jsonify({"error": "仅支持 .xlsx 文件"}), 400

    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        return jsonify({"error": "文件无法解析，请使用模板格式"}), 400

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return jsonify({"error": "文件为空"}), 400

    header = [str(c).strip() if c is not None else "" for c in rows[0]]

    def col(name):
        return header.index(name) if name in header else -1

    type_col = col("题型")
    title_col = col("题干")
    ans_col = col("正确答案")
    score_col = col("分值")
    opt_cols = []
    opt_letters = []
    for i, h in enumerate(header):
        if h.startswith("选项") and h[2:] and h[2:] in string.ascii_uppercase:
            opt_cols.append(i)
            opt_letters.append(h[2:])

    if type_col < 0 or title_col < 0 or ans_col < 0:
        return jsonify({"error": "缺少必要列：题型 / 题干 / 正确答案"}), 400

    questions = []
    warnings = []
    for idx, row in enumerate(rows[1:], start=2):
        if row is None:
            continue
        cells = list(row) + [None] * 8
        type_raw = str(cells[type_col]).strip() if cells[type_col] is not None else ""
        title = str(cells[title_col]).strip() if cells[title_col] is not None else ""
        if not type_raw and not title:
            continue
        if type_raw not in EXAM_TYPE_MAP:
            warnings.append(f"第{idx}行：未知题型「{type_raw}」，已跳过")
            continue
        if not title:
            warnings.append(f"第{idx}行：缺少题干，已跳过")
            continue

        qtype = EXAM_TYPE_MAP[type_raw]

        opt_pairs = []
        for cidx, letter in zip(opt_cols, opt_letters):
            if cidx >= len(cells):
                continue
            text = str(cells[cidx]).strip() if cells[cidx] is not None else ""
            if text:
                opt_pairs.append((letter, text))
        options = [text for _, text in opt_pairs]

        answer = str(cells[ans_col]).strip() if cells[ans_col] is not None else ""
        score_val = cells[score_col] if score_col >= 0 else None
        try:
            score = int(score_val) if score_val not in (None, "") else 0
        except (ValueError, TypeError):
            score = 0

        if qtype == "judge":
            letter = answer.upper()
            if letter and letter not in ("A", "B"):
                warnings.append(f"第{idx}行：判断题正确答案应为 A（对）或 B（错）")
                letter = "A"
            correct = letter == "A"
            q = {
                "type": "judge", "title": title, "score": score,
                "options": [
                    {"text": "对", "is_correct": correct},
                    {"text": "错", "is_correct": not correct},
                ],
            }
        elif qtype == "fill":
            ref = str(cells[ans_col]).strip() if cells[ans_col] is not None else ""
            if not ref:
                warnings.append(f"第{idx}行：填空题缺少参考答案")
            q = {"type": "fill", "title": title, "score": score, "answer": ref}
        else:
            if not options:
                warnings.append(f"第{idx}行：缺少选项，已跳过")
                continue
            letters = "".join(ch for ch in answer.upper() if ch in string.ascii_uppercase)
            correct_set = set(letters)
            if not letters:
                warnings.append(f"第{idx}行：缺少正确答案")
            opt_objs = []
            for letter, text in opt_pairs:
                opt_objs.append({"text": text, "is_correct": letter in correct_set})
            q = {"type": qtype, "title": title, "score": score, "options": opt_objs}

        questions.append(q)

    if not questions:
        first_warn = warnings[0] if warnings else ""
        return jsonify({"error": "未解析到任何题目" + (f"（{first_warn}）" if first_warn else "")}), 400

    return jsonify({"questions": questions, "warnings": warnings})


@app.route("/api/me/ai-key", methods=["GET"])
def get_ai_key():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401
    return jsonify({"has_key": bool((user["ai_key"] or "").strip())})


@app.route("/api/me/ai-key", methods=["POST"])
def set_ai_key():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401
    data = request.get_json(silent=True) or {}
    ai_key = (data.get("ai_key") or "").strip()
    conn = get_db()
    conn.execute("UPDATE users SET ai_key = ? WHERE id = ?", (ai_key, user["id"]))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/me/ai-key", methods=["DELETE"])
def delete_ai_key():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401
    conn = get_db()
    conn.execute("UPDATE users SET ai_key = '' WHERE id = ?", (user["id"],))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/ai/generate", methods=["POST"])
def ai_generate():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401
    data = request.get_json(silent=True) or {}
    kind = data.get("type") or ""
    prompt = (data.get("prompt") or "").strip()
    if not prompt:
        return jsonify({"error": "请填写出卷描述"}), 400
    api_key = (user["ai_key"] or "").strip()
    if not api_key:
        return jsonify({"error": "请先在 AI 设置里填写你的 DeepSeek API Key"}), 400
    mode = data.get("mode") or "score"
    if kind == "exam":
        system = EXAM_AI_SYSTEM
    elif kind == "test":
        system = JUMP_AI_SYSTEM if mode == "jump" else TEST_AI_SYSTEM
    else:
        return jsonify({"error": "不支持的问卷类型"}), 400

    try:
        result = call_deepseek(api_key, system, prompt)
    except HTTPError as e:
        if e.code in (401, 403):
            return jsonify({"error": "API Key 无效或没有权限，请检查后重试"}), 400
        return jsonify({"error": f"AI 服务返回错误（HTTP {e.code}），请稍后重试"}), 502
    except Exception as e:
        return jsonify({"error": "AI 调用失败：" + str(e)}), 502

    if not isinstance(result, dict):
        return jsonify({"error": "AI 返回格式异常，请重试"}), 502
    questions = result.get("questions")
    if not isinstance(questions, list):
        return jsonify({"error": "AI 返回缺少题目，请重试"}), 502
    out = {"questions": questions}
    if kind == "test":
        cards = result.get("result_cards")
        out["result_cards"] = cards if isinstance(cards, list) else []
        if mode == "jump":
            sanitize_jump(out["questions"], out["result_cards"])
    return jsonify(out)


@app.route("/api/me/email/send-code", methods=["POST"])
def send_email_code():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    if not config.SMTP_USER or not config.SMTP_PASSWORD:
        return jsonify({"error": "SMTP 未配置，请在 config.py 中填写邮箱和授权码"}), 500

    data = request.get_json(silent=True) or {}
    email = (data.get("email") or "").strip()
    if not email or "@" not in email:
        return jsonify({"error": "邮箱格式不正确"}), 400

    conn = get_db()
    exists = conn.execute(
        "SELECT id FROM users WHERE email = ? AND id != ?", (email, user["id"])
    ).fetchone()
    if exists:
        conn.close()
        return jsonify({"error": "该邮箱已被绑定"}), 400

    code = f"{secrets.randbelow(1000000):06d}"
    conn.execute("DELETE FROM email_codes WHERE user_id = ?", (user["id"],))
    conn.execute(
        "INSERT INTO email_codes (user_id, email, code, expires_at) "
        "VALUES (?, ?, ?, datetime('now', '+10 minutes'))",
        (user["id"], email, code),
    )
    conn.commit()
    conn.close()

    try:
        send_verification_email(email, code)
    except Exception as e:
        return jsonify({"error": f"邮件发送失败：{e}"}), 500

    return jsonify({"ok": True})


@app.route("/api/me/email/verify", methods=["POST"])
def verify_email_code():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    data = request.get_json(silent=True) or {}
    email = (data.get("email") or "").strip()
    code = (data.get("code") or "").strip()

    conn = get_db()
    row = conn.execute(
        "SELECT * FROM email_codes WHERE user_id = ? AND email = ? "
        "AND expires_at > datetime('now') ORDER BY id DESC LIMIT 1",
        (user["id"], email),
    ).fetchone()
    if row is None:
        conn.close()
        return jsonify({"error": "验证码已过期，请重新发送"}), 400
    if row["code"] != code:
        conn.close()
        return jsonify({"error": "验证码错误"}), 400

    exists = conn.execute(
        "SELECT id FROM users WHERE email = ? AND id != ?", (email, user["id"])
    ).fetchone()
    if exists:
        conn.close()
        return jsonify({"error": "该邮箱已被绑定"}), 400

    conn.execute("UPDATE users SET email = ? WHERE id = ?", (email, user["id"]))
    conn.execute("DELETE FROM email_codes WHERE user_id = ?", (user["id"],))
    conn.commit()
    conn.close()

    return jsonify({"ok": True, "email": email})


@app.route("/api/password/forgot", methods=["POST"])
def forgot_password():
    if not config.SMTP_USER or not config.SMTP_PASSWORD:
        return jsonify({"error": "SMTP 未配置，请在 config.py 中填写邮箱和授权码"}), 500

    data = request.get_json(silent=True) or {}
    email = (data.get("email") or "").strip()
    if not email or "@" not in email:
        return jsonify({"error": "邮箱格式不正确"}), 400

    conn = get_db()
    user = conn.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if user is None:
        conn.close()
        return jsonify({"error": "该邮箱未绑定账号"}), 400

    code = f"{secrets.randbelow(1000000):06d}"
    conn.execute("DELETE FROM email_codes WHERE user_id = ?", (user["id"],))
    conn.execute(
        "INSERT INTO email_codes (user_id, email, code, expires_at) "
        "VALUES (?, ?, ?, datetime('now', '+10 minutes'))",
        (user["id"], email, code),
    )
    conn.commit()
    conn.close()

    try:
        send_verification_email(email, code)
    except Exception as e:
        return jsonify({"error": f"邮件发送失败：{e}"}), 500

    return jsonify({"ok": True})


@app.route("/api/password/reset", methods=["POST"])
def reset_password():
    data = request.get_json(silent=True) or {}
    email = (data.get("email") or "").strip()
    code = (data.get("code") or "").strip()
    new_password = data.get("new_password") or ""

    if not new_password:
        return jsonify({"error": "新密码不能为空"}), 400
    err = validate_password(new_password)
    if err:
        return jsonify({"error": err}), 400

    conn = get_db()
    user = conn.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if user is None:
        conn.close()
        return jsonify({"error": "该邮箱未绑定账号"}), 400

    if check_password_hash(user["password_hash"], new_password):
        conn.close()
        return jsonify({"error": "新密码不能与旧密码相同"}), 400

    row = conn.execute(
        "SELECT * FROM email_codes WHERE user_id = ? AND email = ? "
        "AND expires_at > datetime('now') ORDER BY id DESC LIMIT 1",
        (user["id"], email),
    ).fetchone()
    if row is None:
        conn.close()
        return jsonify({"error": "验证码已过期，请重新发送"}), 400
    if row["code"] != code:
        conn.close()
        return jsonify({"error": "验证码错误"}), 400

    conn.execute(
        "UPDATE users SET password_hash = ? WHERE id = ?",
        (generate_password_hash(new_password), user["id"]),
    )
    conn.execute("DELETE FROM email_codes WHERE user_id = ?", (user["id"],))
    conn.commit()
    conn.close()

    return jsonify({"ok": True})


@app.route("/api/questionnaires", methods=["POST"])
def create_questionnaire():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    data = request.get_json(silent=True) or {}
    qtype = data.get("type") or ""
    title = (data.get("title") or "").strip()
    description = (data.get("description") or "").strip()
    result_mode = data.get("result_mode") or "score"
    if not title:
        return jsonify({"error": "问卷名不能为空"}), 400

    conn = get_db()
    cur = conn.execute(
        "INSERT INTO questionnaires (user_id, type, title, description, result_mode) VALUES (?, ?, ?, ?, ?)",
        (user["id"], qtype, title, description, result_mode),
    )
    conn.commit()
    qid = cur.lastrowid
    conn.close()

    return jsonify({"id": qid, "type": qtype, "title": title, "description": description, "status": "draft", "result_mode": result_mode})


@app.route("/api/questionnaires/mine", methods=["GET"])
def list_mine():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    rows = conn.execute(
        "SELECT id, type, title, status, updated_at, published_at, "
        "(SELECT COUNT(*) FROM responses WHERE questionnaire_id = questionnaires.id) AS response_count "
        "FROM questionnaires "
        "WHERE user_id = ? ORDER BY updated_at DESC",
        (user["id"],),
    ).fetchall()
    conn.close()

    return jsonify([dict(r) for r in rows])


@app.route("/api/questionnaires/<int:qid>", methods=["GET"])
def get_questionnaire(qid):
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    q = conn.execute("SELECT * FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404

    questions = conn.execute(
        "SELECT * FROM questions WHERE questionnaire_id = ? ORDER BY order_index",
        (qid,),
    ).fetchall()
    question_list = []
    for question in questions:
        opts = conn.execute(
            "SELECT text, score, is_correct, jump_to FROM options WHERE question_id = ? ORDER BY order_index",
            (question["id"],),
        ).fetchall()
        question_list.append({
            "type": question["type"],
            "title": question["title"],
            "score": question["score"],
            "answer": question["answer"],
            "options": [
                {"text": o["text"], "score": o["score"], "is_correct": bool(o["is_correct"]), "jump_to": o["jump_to"] or ""}
                for o in opts
            ],
        })

    cards = conn.execute(
        "SELECT min_score, max_score, text, label FROM result_cards "
        "WHERE questionnaire_id = ? ORDER BY order_index",
        (qid,),
    ).fetchall()
    conn.close()

    return jsonify({
        "id": q["id"],
        "type": q["type"],
        "title": q["title"],
        "description": q["description"],
        "status": q["status"],
        "full_score": q["full_score"],
        "result_mode": q["result_mode"],
        "questions": question_list,
        "result_cards": [dict(c) for c in cards],
    })


@app.route("/api/questionnaires/<int:qid>", methods=["PUT"])
def update_questionnaire(qid):
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    try:
        q = conn.execute("SELECT * FROM questionnaires WHERE id = ?", (qid,)).fetchone()
        if q is None:
            return jsonify({"error": "问卷不存在"}), 404
        if q["user_id"] != user["id"]:
            return jsonify({"error": "无权限"}), 403

        data = request.get_json(silent=True) or {}
        title = (data.get("title") or "").strip()
        description = (data.get("description") or "").strip()
        questions = data.get("questions") or []
        result_cards = data.get("result_cards") or []
        status = data.get("status")
        full_score = data.get("full_score")
        if full_score in (None, ""):
            full_score = None
        else:
            full_score = int(full_score)
        result_mode = data.get("result_mode") or "score"

        if status in ("draft", "published", "stopped"):
            if status == "published":
                hits = find_blocked_words(
                    title,
                    description,
                    *[qq.get("title") for qq in questions],
                    *[o.get("text") for qq in questions for o in (qq.get("options") or [])],
                )
                if hits:
                    return jsonify({"error": "内容包含敏感词，无法发布：" + "、".join(hits)}), 400
                conn.execute(
                    "UPDATE questionnaires SET title = ?, description = ?, status = ?, full_score = ?, result_mode = ?, updated_at = datetime('now'), published_at = datetime('now') WHERE id = ?",
                    (title, description, status, full_score, result_mode, qid),
                )
            else:
                conn.execute(
                    "UPDATE questionnaires SET title = ?, description = ?, status = ?, full_score = ?, result_mode = ?, updated_at = datetime('now') WHERE id = ?",
                    (title, description, status, full_score, result_mode, qid),
                )
        else:
            conn.execute(
                "UPDATE questionnaires SET title = ?, description = ?, full_score = ?, result_mode = ?, updated_at = datetime('now') WHERE id = ?",
                (title, description, full_score, result_mode, qid),
            )
        conn.execute(
            "DELETE FROM options WHERE question_id IN "
            "(SELECT id FROM questions WHERE questionnaire_id = ?)",
            (qid,),
        )
        conn.execute("DELETE FROM questions WHERE questionnaire_id = ?", (qid,))
        conn.execute("DELETE FROM result_cards WHERE questionnaire_id = ?", (qid,))

        for qi, qq in enumerate(questions):
            cur = conn.execute(
                "INSERT INTO questions (questionnaire_id, type, title, order_index, score, answer) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (qid, qq.get("type"), qq.get("title"), qi, qq.get("score") or 0, qq.get("answer") or ""),
            )
            question_id = cur.lastrowid
            for oi, o in enumerate(qq.get("options") or []):
                conn.execute(
                    "INSERT INTO options (question_id, text, score, order_index, is_correct, jump_to) "
                    "VALUES (?, ?, ?, ?, ?, ?)",
                    (question_id, o.get("text"), o.get("score") or 0, oi, 1 if o.get("is_correct") else 0, o.get("jump_to") or ""),
                )

        for ci, c in enumerate(result_cards):
            conn.execute(
                "INSERT INTO result_cards (questionnaire_id, min_score, max_score, text, order_index, label) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (qid, c.get("min_score") or 0, c.get("max_score") or 0, c.get("text"), ci, c.get("label") or ""),
            )

        conn.commit()
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/questionnaires/<int:qid>/status", methods=["POST"])
def set_questionnaire_status(qid):
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    q = conn.execute("SELECT * FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404
    if q["user_id"] != user["id"]:
        conn.close()
        return jsonify({"error": "无权限"}), 403

    data = request.get_json(silent=True) or {}
    status = data.get("status")
    if status == "published":
        hits = scan_questionnaire(qid)
        if hits:
            conn.close()
            return jsonify({"error": "内容包含敏感词，无法发布：" + "、".join(hits)}), 400
        conn.execute(
            "UPDATE questionnaires SET status = 'published', updated_at = datetime('now'), published_at = datetime('now') WHERE id = ?",
            (qid,),
        )
    elif status == "stopped":
        conn.execute(
            "UPDATE questionnaires SET status = 'stopped', updated_at = datetime('now') WHERE id = ?",
            (qid,),
        )
    else:
        conn.close()
        return jsonify({"error": "无效状态"}), 400

    conn.commit()
    conn.close()
    return jsonify({"ok": True, "status": status})


@app.route("/api/questionnaires/<int:qid>", methods=["DELETE"])
def delete_questionnaire(qid):
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    q = conn.execute("SELECT * FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404
    if q["user_id"] != user["id"]:
        conn.close()
        return jsonify({"error": "无权限"}), 403

    conn.execute(
        "DELETE FROM options WHERE question_id IN "
        "(SELECT id FROM questions WHERE questionnaire_id = ?)",
        (qid,),
    )
    conn.execute("DELETE FROM questions WHERE questionnaire_id = ?", (qid,))
    conn.execute("DELETE FROM result_cards WHERE questionnaire_id = ?", (qid,))
    conn.execute("DELETE FROM questionnaires WHERE id = ?", (qid,))
    conn.commit()
    conn.close()

    return jsonify({"ok": True})


@app.route("/api/questionnaires/public", methods=["GET"])
def public_questionnaires():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    rows = conn.execute(
        "SELECT q.id, q.title, q.type, q.description, q.status, q.updated_at, q.user_id, u.username AS author, "
        "(SELECT COUNT(*) FROM user_actions WHERE questionnaire_id = q.id AND action = 'favorite') AS favorite_count, "
        "(SELECT COUNT(*) FROM user_actions WHERE questionnaire_id = q.id AND action = 'like') AS like_count, "
        "EXISTS(SELECT 1 FROM user_actions WHERE user_id = ? AND questionnaire_id = q.id AND action = 'favorite') AS favorited, "
        "EXISTS(SELECT 1 FROM user_actions WHERE user_id = ? AND questionnaire_id = q.id AND action = 'like') AS liked "
        "FROM questionnaires q JOIN users u ON u.id = q.user_id "
        "WHERE q.status = 'published' ORDER BY q.updated_at DESC",
        (user["id"], user["id"]),
    ).fetchall()
    conn.close()

    result = []
    for r in rows:
        d = dict(r)
        d["favorited"] = bool(d["favorited"])
        d["liked"] = bool(d["liked"])
        d["is_owner"] = d["user_id"] == user["id"]
        d.pop("user_id", None)
        result.append(d)
    return jsonify(result)


@app.route("/api/questionnaires/<int:qid>/public", methods=["GET"])
def public_questionnaire(qid):
    user = current_user()
    conn = get_db()
    q = conn.execute(
        "SELECT q.id, q.type, q.title, q.description, q.status, q.result_mode, q.user_id, u.username AS author "
        "FROM questionnaires q JOIN users u ON u.id = q.user_id WHERE q.id = ?",
        (qid,),
    ).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404
    is_owner = user is not None and q["user_id"] == user["id"]
    if q["status"] != "published" and not is_owner:
        conn.close()
        return jsonify({"error": "问卷未发布"}), 404

    questions = conn.execute(
        "SELECT * FROM questions WHERE questionnaire_id = ? ORDER BY order_index",
        (qid,),
    ).fetchall()
    question_list = []
    for question in questions:
        opts = conn.execute(
            "SELECT text, jump_to FROM options WHERE question_id = ? ORDER BY order_index",
            (question["id"],),
        ).fetchall()
        question_list.append({
            "type": question["type"],
            "title": question["title"],
            "options": [{"text": o["text"], "jump_to": o["jump_to"] or ""} for o in opts],
        })
    conn.close()

    return jsonify({
        "id": q["id"],
        "type": q["type"],
        "title": q["title"],
        "description": q["description"],
        "author": q["author"],
        "status": q["status"],
        "result_mode": q["result_mode"],
        "is_owner": is_owner,
        "questions": question_list,
    })


@app.route("/api/questionnaires/<int:qid>/action", methods=["POST"])
def questionnaire_action(qid):
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    data = request.get_json(silent=True) or {}
    action = data.get("action")
    if action not in ("favorite", "like", "browse"):
        return jsonify({"error": "无效操作"}), 400

    conn = get_db()
    q = conn.execute("SELECT id FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404

    existing = conn.execute(
        "SELECT id FROM user_actions WHERE user_id = ? AND questionnaire_id = ? AND action = ?",
        (user["id"], qid, action),
    ).fetchone()

    if action in ("favorite", "like"):
        if existing:
            conn.execute("DELETE FROM user_actions WHERE id = ?", (existing["id"],))
            active = False
        else:
            conn.execute(
                "INSERT INTO user_actions (user_id, questionnaire_id, action) VALUES (?, ?, ?)",
                (user["id"], qid, action),
            )
            active = True
        conn.commit()
        conn.close()
        return jsonify({"active": active})

    if not existing:
        conn.execute(
            "INSERT INTO user_actions (user_id, questionnaire_id, action) VALUES (?, ?, ?)",
            (user["id"], qid, action),
        )
        conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/me/actions", methods=["GET"])
def my_actions():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    rows = conn.execute(
        "SELECT ua.action, q.id, q.title, q.type, q.status, q.updated_at "
        "FROM user_actions ua JOIN questionnaires q ON q.id = ua.questionnaire_id "
        "WHERE ua.user_id = ? ORDER BY ua.created_at DESC",
        (user["id"],),
    ).fetchall()
    conn.close()

    return jsonify([dict(r) for r in rows])


@app.route("/api/me/notifications", methods=["GET"])
def my_notifications():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM ("
        "SELECT 'submit' AS kind, r.created_at, u.username AS actor, q.id AS qid, q.title "
        "FROM responses r "
        "JOIN questionnaires q ON q.id = r.questionnaire_id "
        "LEFT JOIN users u ON u.id = r.user_id "
        "WHERE q.user_id = ? "
        "UNION ALL "
        "SELECT ua.action AS kind, ua.created_at, u2.username AS actor, q.id AS qid, q.title "
        "FROM user_actions ua "
        "JOIN questionnaires q ON q.id = ua.questionnaire_id "
        "JOIN users u2 ON u2.id = ua.user_id "
        "WHERE q.user_id = ? AND ua.action IN ('like', 'favorite')"
        ") ORDER BY created_at DESC",
        (user["id"], user["id"]),
    ).fetchall()
    conn.close()

    return jsonify([dict(r) for r in rows])


@app.route("/api/me/responses", methods=["GET"])
def my_responses():
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    rows = conn.execute(
        "SELECT q.id, q.title, q.type, q.user_id AS author_id, q.full_score "
        "FROM responses r JOIN questionnaires q ON q.id = r.questionnaire_id "
        "WHERE r.user_id = ? GROUP BY q.id ORDER BY MAX(r.created_at) DESC",
        (user["id"],),
    ).fetchall()

    result = []
    for row in rows:
        resp_rows = conn.execute(
            "SELECT id, total_score, result_text, answers, created_at FROM responses "
            "WHERE questionnaire_id = ? AND user_id = ? ORDER BY created_at DESC, id DESC",
            (row["id"], user["id"]),
        ).fetchall()
        responses = []
        for rr in resp_rows:
            answers = []
            if rr["answers"]:
                try:
                    answers = json.loads(rr["answers"])
                except Exception:
                    answers = []
            responses.append({
                "id": rr["id"],
                "total_score": rr["total_score"],
                "result_text": rr["result_text"],
                "answers": answers,
                "created_at": rr["created_at"],
            })

        full_score = row["full_score"]
        if full_score is None and row["type"] == "exam":
            full_score = conn.execute(
                "SELECT COALESCE(SUM(score), 0) FROM questions WHERE questionnaire_id = ?",
                (row["id"],),
            ).fetchone()[0]

        result.append({
            "id": row["id"],
            "title": row["title"],
            "type": row["type"],
            "mine": row["author_id"] == user["id"],
            "full_score": full_score,
            "responses": responses,
        })

    conn.close()
    return jsonify(result)


@app.route("/api/questionnaires/<int:qid>/stats", methods=["GET"])
def questionnaire_stats(qid):
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    q = conn.execute("SELECT * FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404
    if q["user_id"] != user["id"]:
        conn.close()
        return jsonify({"error": "无权查看"}), 403

    responses = conn.execute(
        "SELECT total_score, result_text FROM responses WHERE questionnaire_id = ?",
        (qid,),
    ).fetchall()
    total = len(responses)

    if q["type"] == "exam":
        ranges = [
            ("0-40%", 0, 40),
            ("41-60%", 41, 60),
            ("61-80%", 61, 80),
            ("81-100%", 81, 100),
        ]
        full = q["full_score"] or 0
        counts = [0, 0, 0, 0]
        for r in responses:
            pct = round((r["total_score"] or 0) / full * 100) if full else 0
            if pct <= 40:
                counts[0] += 1
            elif pct <= 60:
                counts[1] += 1
            elif pct <= 80:
                counts[2] += 1
            else:
                counts[3] += 1
        items = [{"label": ranges[i][0], "count": counts[i]} for i in range(4)]
    else:
        cards = conn.execute(
            "SELECT min_score, max_score, text, label FROM result_cards "
            "WHERE questionnaire_id = ? ORDER BY order_index",
            (qid,),
        ).fetchall()

        if q["result_mode"] == "jump":
            text_to_label = {c["text"]: c["label"] or "" for c in cards}
            counts = {}
            for r in responses:
                lbl = text_to_label.get(r["result_text"], "")
                key = ("结果 " + lbl) if lbl else (r["result_text"] or "未匹配")
                counts[key] = counts.get(key, 0) + 1
            items = [{"label": k, "count": v} for k, v in counts.items()]
        else:
            matched = set()
            items = []
            for c in cards:
                lo, hi = c["min_score"], c["max_score"]
                cnt = 0
                for ri, r in enumerate(responses):
                    if lo <= (r["total_score"] or 0) <= hi:
                        cnt += 1
                        matched.add(ri)
                label = c["label"] or f"{lo}-{hi}分"
                items.append({"label": label, "count": cnt})
            leftover = sum(1 for ri in range(total) if ri not in matched)
            if leftover:
                items.append({"label": "未匹配", "count": leftover})

    conn.close()
    return jsonify({"type": q["type"], "items": items, "total": total})


@app.route("/api/responses/<int:rid>", methods=["DELETE"])
def delete_response(rid):
    user = current_user()
    if not user:
        return jsonify({"error": "未登录"}), 401

    conn = get_db()
    r = conn.execute(
        "SELECT id FROM responses WHERE id = ? AND user_id = ?", (rid, user["id"])
    ).fetchone()
    if r is None:
        conn.close()
        return jsonify({"error": "记录不存在"}), 404

    conn.execute("DELETE FROM responses WHERE id = ?", (rid,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/questionnaires/<int:qid>/submit", methods=["POST"])
def submit_questionnaire(qid):
    data = request.get_json(silent=True) or {}
    answers = data.get("answers") or []
    user = current_user()
    user_id = user["id"] if user else None

    conn = get_db()
    q = conn.execute("SELECT * FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404
    if q["status"] != "published":
        conn.close()
        return jsonify({"error": "问卷未发布"}), 404

    questions = conn.execute(
        "SELECT * FROM questions WHERE questionnaire_id = ? ORDER BY order_index",
        (qid,),
    ).fetchall()

    if q["type"] == "exam":
        total_score = 0
        answer_detail = []
        review = []
        for i, question in enumerate(questions):
            qtype = question["type"]
            qscore = question["score"] or 0
            ans = answers[i] if i < len(answers) else None
            got = 0
            opts = conn.execute(
                "SELECT text, is_correct FROM options WHERE question_id = ? ORDER BY order_index",
                (question["id"],),
            ).fetchall()

            correct_texts = [o["text"] for o in opts if o["is_correct"]]
            my_texts = []
            opt_list = []

            if qtype in ("single", "judge"):
                if isinstance(ans, int) and 0 <= ans < len(opts) and opts[ans]["is_correct"]:
                    got = qscore
                if isinstance(ans, int) and 0 <= ans < len(opts):
                    my_texts = [opts[ans]["text"]]
                opt_list = [
                    {"text": o["text"], "is_correct": bool(o["is_correct"]), "chosen": ans == oi}
                    for oi, o in enumerate(opts)
                ]
            elif qtype == "multiple":
                if isinstance(ans, list):
                    correct = {oi for oi, o in enumerate(opts) if o["is_correct"]}
                    chosen = {oi for oi in ans if isinstance(oi, int)}
                    if chosen == correct:
                        got = qscore
                if isinstance(ans, list):
                    my_texts = [opts[oi]["text"] for oi in ans if isinstance(oi, int) and 0 <= oi < len(opts)]
                opt_list = [
                    {"text": o["text"], "is_correct": bool(o["is_correct"]), "chosen": (isinstance(ans, list) and oi in ans)}
                    for oi, o in enumerate(opts)
                ]
            elif qtype == "fill":
                refs = [a.strip() for a in (question["answer"] or "").split("|") if a.strip()]
                if isinstance(ans, str) and ans.strip() in refs:
                    got = qscore
                correct_texts = refs
                if isinstance(ans, str) and ans.strip():
                    my_texts = [ans.strip()]

            total_score += got
            answer_detail.append({
                "title": question["title"],
                "answer": ans,
                "correct": got > 0,
                "score": got,
            })
            review.append({
                "title": question["title"],
                "type": qtype,
                "score": got,
                "correct": got > 0,
                "options": opt_list,
                "correct_answer": correct_texts,
                "my_answer": my_texts,
            })

        sum_score = sum((qq["score"] or 0) for qq in questions)
        full_score = q["full_score"] if q["full_score"] else sum_score

        conn.execute(
            "INSERT INTO responses (questionnaire_id, user_id, answers, total_score, result_text) VALUES (?, ?, ?, ?, ?)",
            (qid, user_id, json.dumps(answer_detail, ensure_ascii=False), total_score, ""),
        )

        if user_id is not None:
            existing = conn.execute(
                "SELECT id FROM user_actions WHERE user_id = ? AND questionnaire_id = ? AND action = 'use'",
                (user_id, qid),
            ).fetchone()
            if not existing:
                conn.execute(
                    "INSERT INTO user_actions (user_id, questionnaire_id, action) VALUES (?, ?, 'use')",
                    (user_id, qid),
                )

        conn.commit()
        conn.close()

        return jsonify({
            "type": "exam",
            "total_score": total_score,
            "full_score": full_score,
            "review": review,
        })

    if q["result_mode"] == "jump":
        result_label = data.get("result_label") or ""
        answer_detail = []
        for i, question in enumerate(questions):
            opts = conn.execute(
                "SELECT text FROM options WHERE question_id = ? ORDER BY order_index",
                (question["id"],),
            ).fetchall()
            ans = answers[i] if i < len(answers) else None
            selected = None
            if isinstance(ans, int) and 0 <= ans < len(opts):
                selected = opts[ans]["text"]
            answer_detail.append({
                "title": question["title"],
                "selected": selected,
            })
        card = conn.execute(
            "SELECT text FROM result_cards WHERE questionnaire_id = ? AND label = ?",
            (qid, result_label),
        ).fetchone()
        result_text = card["text"] if card else "暂无匹配的结果"

        conn.execute(
            "INSERT INTO responses (questionnaire_id, user_id, answers, total_score, result_text) VALUES (?, ?, ?, ?, ?)",
            (qid, user_id, json.dumps(answer_detail, ensure_ascii=False), 0, result_text),
        )

        if user_id is not None:
            existing = conn.execute(
                "SELECT id FROM user_actions WHERE user_id = ? AND questionnaire_id = ? AND action = 'use'",
                (user_id, qid),
            ).fetchone()
            if not existing:
                conn.execute(
                    "INSERT INTO user_actions (user_id, questionnaire_id, action) VALUES (?, ?, 'use')",
                    (user_id, qid),
                )

        conn.commit()
        conn.close()

        return jsonify({
            "label": result_label,
            "result_text": result_text,
        })

    total_score = 0
    answer_detail = []
    for i, question in enumerate(questions):
        opts = conn.execute(
            "SELECT text, score FROM options WHERE question_id = ? ORDER BY order_index",
            (question["id"],),
        ).fetchall()
        oi = answers[i] if i < len(answers) else None
        selected = None
        if oi is not None and 0 <= oi < len(opts):
            selected = opts[oi]
        score = selected["score"] if selected else 0
        total_score += score
        answer_detail.append({
            "title": question["title"],
            "selected": selected["text"] if selected else None,
            "score": score,
        })

    cards = conn.execute(
        "SELECT min_score, max_score, text FROM result_cards WHERE questionnaire_id = ? ORDER BY order_index",
        (qid,),
    ).fetchall()
    result_text = None
    for c in cards:
        if c["min_score"] <= total_score <= c["max_score"]:
            result_text = c["text"]
            break
    if result_text is None:
        result_text = "暂无匹配的结果"

    conn.execute(
        "INSERT INTO responses (questionnaire_id, user_id, answers, total_score, result_text) VALUES (?, ?, ?, ?, ?)",
        (qid, user_id, json.dumps(answer_detail, ensure_ascii=False), total_score, result_text),
    )

    if user_id is not None:
        existing = conn.execute(
            "SELECT id FROM user_actions WHERE user_id = ? AND questionnaire_id = ? AND action = 'use'",
            (user_id, qid),
        ).fetchone()
        if not existing:
            conn.execute(
                "INSERT INTO user_actions (user_id, questionnaire_id, action) VALUES (?, ?, 'use')",
                (user_id, qid),
            )

    conn.commit()
    conn.close()

    return jsonify({
        "total_score": total_score,
        "result_text": result_text,
    })


@app.route("/api/admin/users", methods=["GET"])
def admin_users():
    user = current_user()
    if not is_super(user):
        return jsonify({"error": "无权限"}), 403
    conn = get_db()
    rows = conn.execute(
        "SELECT id, username, role, email, created_at FROM users ORDER BY id"
    ).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route("/api/admin/users/<int:uid>/role", methods=["POST"])
def admin_set_role(uid):
    user = current_user()
    if not is_super(user):
        return jsonify({"error": "无权限"}), 403
    data = request.get_json(silent=True) or {}
    role = data.get("role")
    if role not in ("admin", "user"):
        return jsonify({"error": "无效角色"}), 400
    conn = get_db()
    target = conn.execute("SELECT * FROM users WHERE id = ?", (uid,)).fetchone()
    if target is None:
        conn.close()
        return jsonify({"error": "用户不存在"}), 404
    if target["role"] == "super":
        conn.close()
        return jsonify({"error": "不能修改最高管理员"}), 403
    conn.execute("UPDATE users SET role = ? WHERE id = ?", (role, uid))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/admin/questionnaires", methods=["GET"])
def admin_questionnaires():
    user = current_user()
    if not is_admin(user):
        return jsonify({"error": "无权限"}), 403
    conn = get_db()
    rows = conn.execute(
        "SELECT q.id, q.title, q.type, q.status, q.updated_at, u.username AS author, "
        "(SELECT COUNT(*) FROM responses WHERE questionnaire_id = q.id) AS response_count "
        "FROM questionnaires q JOIN users u ON u.id = q.user_id ORDER BY q.updated_at DESC"
    ).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route("/api/admin/questionnaires/<int:qid>/status", methods=["POST"])
def admin_set_questionnaire_status(qid):
    user = current_user()
    if not is_admin(user):
        return jsonify({"error": "无权限"}), 403
    data = request.get_json(silent=True) or {}
    status = data.get("status")
    if status not in ("published", "stopped"):
        return jsonify({"error": "无效状态"}), 400
    conn = get_db()
    q = conn.execute("SELECT * FROM questionnaires WHERE id = ?", (qid,)).fetchone()
    if q is None:
        conn.close()
        return jsonify({"error": "问卷不存在"}), 404
    if status == "published":
        conn.execute(
            "UPDATE questionnaires SET status = 'published', updated_at = datetime('now'), published_at = datetime('now') WHERE id = ?",
            (qid,),
        )
    else:
        conn.execute(
            "UPDATE questionnaires SET status = 'stopped', updated_at = datetime('now') WHERE id = ?",
            (qid,),
        )
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "status": status})


# 前端打包产物目录（生产环境由后端统一托管）
DIST_DIR = Path(__file__).parent.parent / "frontend" / "dist"


@app.route("/")
def serve_index():
    return send_from_directory(DIST_DIR, "index.html")


@app.route("/<path:path>")
def serve_frontend(path):
    if path.startswith("api/") or path.startswith("uploads/"):
        return jsonify({"error": "not found"}), 404
    candidate = DIST_DIR / path
    if candidate.is_file():
        return send_from_directory(DIST_DIR, path)
    return send_from_directory(DIST_DIR, "index.html")


if __name__ == "__main__":
    import os

    if os.environ.get("FLASK_DEV") == "1":
        # 本地开发模式：Flask 自带服务器 + 热重载，跑在 5000（前端 Vite 代理指向这里）
        app.run(debug=True, host="0.0.0.0", port=5000)
    else:
        # 生产模式：waitress，端口由 PORT 环境变量决定（默认 8080）
        from waitress import serve

        port = int(os.environ.get("PORT", "8080"))
        print(f"问卷系统已启动：http://0.0.0.0:{port}")
        serve(app, host="0.0.0.0", port=port)
